"""
Export worker tasks.

Tasks
-----
export_to_csv(query_params, output_path)
    Query businesses matching *query_params* and write a CSV file.

export_to_excel(query_params, output_path)
    Query businesses and write a multi-sheet XLSX file.

export_to_pdf(query_params, output_path)
    Query businesses and write a formatted PDF summary report.

generate_audit_pdf(audit_id, output_path)
    Generate a detailed PDF report for a single AuditReport.

cleanup_old_exports()
    Periodic task: delete export files older than 24 hours.
"""
from __future__ import annotations

import asyncio
import csv
import io
import os
import uuid
from datetime import datetime, timedelta, timezone
from pathlib import Path
from typing import Any, Dict, List, Optional

from celery.utils.log import get_task_logger
from sqlalchemy import select

from app.core.config import settings
from app.core.database import AsyncSessionLocal
from app.models.audit import AuditReport
from app.models.business import Business
from app.workers.celery_app import celery_app

logger = get_task_logger(__name__)

# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------


def _run_async(coro: Any) -> Any:
    loop = asyncio.new_event_loop()
    try:
        return loop.run_until_complete(coro)
    finally:
        loop.close()


def _ensure_dir(path: str) -> Path:
    p = Path(path)
    p.parent.mkdir(parents=True, exist_ok=True)
    return p


# Business columns included in all export formats
_BUSINESS_COLUMNS = [
    "id",
    "name",
    "city",
    "state",
    "country",
    "phone",
    "email",
    "website_url",
    "has_website",
    "website_status",
    "industry",
    "category",
    "source",
    "overall_score",
    "opportunity_score",
    "audit_count",
    "discovered_at",
    "last_audited_at",
    "created_at",
]


async def _query_businesses(query_params: Dict[str, Any]) -> List[Dict[str, Any]]:
    """
    Query the Business table according to *query_params*.

    Supported keys
    --------------
    city, state, country, industry, category, source, has_website,
    min_score, max_score, ids (list[str]), limit (int, default 10000)
    """
    async with AsyncSessionLocal() as session:
        stmt = select(Business)

        if query_params.get("city"):
            stmt = stmt.where(Business.city == query_params["city"])
        if query_params.get("state"):
            stmt = stmt.where(Business.state == query_params["state"])
        if query_params.get("country"):
            stmt = stmt.where(Business.country == query_params["country"])
        if query_params.get("industry"):
            stmt = stmt.where(Business.industry == query_params["industry"])
        if query_params.get("category"):
            stmt = stmt.where(Business.category == query_params["category"])
        if query_params.get("source"):
            stmt = stmt.where(Business.source == query_params["source"])
        if query_params.get("has_website") is not None:
            stmt = stmt.where(Business.has_website == bool(query_params["has_website"]))
        if query_params.get("min_score") is not None:
            stmt = stmt.where(Business.overall_score >= float(query_params["min_score"]))
        if query_params.get("max_score") is not None:
            stmt = stmt.where(Business.overall_score <= float(query_params["max_score"]))
        if query_params.get("ids"):
            ids = [uuid.UUID(i) for i in query_params["ids"]]
            stmt = stmt.where(Business.id.in_(ids))

        limit: int = int(query_params.get("limit", 10_000))
        stmt = stmt.limit(limit)

        rows = await session.execute(stmt)
        businesses = rows.scalars().all()

        result = []
        for biz in businesses:
            row: Dict[str, Any] = {}
            for col in _BUSINESS_COLUMNS:
                val = getattr(biz, col, None)
                if isinstance(val, uuid.UUID):
                    val = str(val)
                elif isinstance(val, datetime):
                    val = val.isoformat()
                elif hasattr(val, "value"):  # enum
                    val = val.value
                row[col] = val
            result.append(row)

        return result


# ---------------------------------------------------------------------------
# Task: export_to_csv
# ---------------------------------------------------------------------------


@celery_app.task(
    name="app.workers.export_worker.export_to_csv",
    queue="export-queue",
    soft_time_limit=300,
    time_limit=360,
)
def export_to_csv(
    query_params: Dict[str, Any],
    output_path: str,
) -> Dict[str, Any]:
    """
    Export matching businesses to a CSV file at *output_path*.

    Returns
    -------
    dict with ``output_path`` and ``row_count``.
    """
    return _run_async(_export_to_csv_async(query_params, output_path))


async def _export_to_csv_async(
    query_params: Dict[str, Any],
    output_path: str,
) -> Dict[str, Any]:
    businesses = await _query_businesses(query_params)
    out = _ensure_dir(output_path)

    with open(out, "w", newline="", encoding="utf-8") as fh:
        if not businesses:
            fh.write(",".join(_BUSINESS_COLUMNS) + "\n")
            logger.info("export_csv.empty", output_path=output_path)
            return {"output_path": output_path, "row_count": 0}

        writer = csv.DictWriter(fh, fieldnames=_BUSINESS_COLUMNS, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(businesses)

    row_count = len(businesses)
    file_size = out.stat().st_size
    logger.info(
        "export_csv.done",
        output_path=output_path,
        row_count=row_count,
        file_size_bytes=file_size,
    )
    return {"output_path": output_path, "row_count": row_count, "file_size": file_size}


# ---------------------------------------------------------------------------
# Task: export_to_excel
# ---------------------------------------------------------------------------


@celery_app.task(
    name="app.workers.export_worker.export_to_excel",
    queue="export-queue",
    soft_time_limit=300,
    time_limit=360,
)
def export_to_excel(
    query_params: Dict[str, Any],
    output_path: str,
) -> Dict[str, Any]:
    """
    Export matching businesses to a multi-sheet XLSX file.

    Sheets
    ------
    Businesses  – main data table
    Summary     – aggregate statistics
    """
    return _run_async(_export_to_excel_async(query_params, output_path))


async def _export_to_excel_async(
    query_params: Dict[str, Any],
    output_path: str,
) -> Dict[str, Any]:
    import openpyxl  # type: ignore
    from openpyxl.styles import Alignment, Font, PatternFill  # type: ignore
    from openpyxl.utils import get_column_letter  # type: ignore

    businesses = await _query_businesses(query_params)
    out = _ensure_dir(output_path)

    wb = openpyxl.Workbook()

    # ---- Sheet 1: Businesses ----
    ws_biz = wb.active
    ws_biz.title = "Businesses"

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(fill_type="solid", fgColor="2563EB")  # blue-600

    # Write header row
    for col_idx, col_name in enumerate(_BUSINESS_COLUMNS, start=1):
        cell = ws_biz.cell(row=1, column=col_idx, value=col_name.replace("_", " ").title())
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")

    # Write data rows
    for row_idx, biz in enumerate(businesses, start=2):
        for col_idx, col_name in enumerate(_BUSINESS_COLUMNS, start=1):
            ws_biz.cell(row=row_idx, column=col_idx, value=biz.get(col_name))

    # Auto-size columns (cap at 50)
    for col_idx in range(1, len(_BUSINESS_COLUMNS) + 1):
        col_letter = get_column_letter(col_idx)
        max_len = max(
            (
                len(str(ws_biz.cell(row=r, column=col_idx).value or ""))
                for r in range(1, min(ws_biz.max_row + 1, 200))
            ),
            default=10,
        )
        ws_biz.column_dimensions[col_letter].width = min(max_len + 2, 50)

    # ---- Sheet 2: Summary ----
    ws_sum = wb.create_sheet("Summary")
    ws_sum.append(["Metric", "Value"])

    total = len(businesses)
    with_website = sum(1 for b in businesses if b.get("has_website"))
    scored = [b["overall_score"] for b in businesses if b.get("overall_score") is not None]
    avg_score = round(sum(scored) / len(scored), 1) if scored else 0

    ws_sum.append(["Total Businesses", total])
    ws_sum.append(["With Website", with_website])
    ws_sum.append(["Without Website", total - with_website])
    ws_sum.append(["Avg Overall Score", avg_score])
    ws_sum.append(["Generated At", datetime.now(timezone.utc).isoformat()])
    ws_sum.append(["Query Params", str(query_params)])

    wb.save(str(out))
    file_size = out.stat().st_size
    logger.info(
        "export_excel.done",
        output_path=output_path,
        row_count=total,
        file_size_bytes=file_size,
    )
    return {"output_path": output_path, "row_count": total, "file_size": file_size}


# ---------------------------------------------------------------------------
# Task: export_to_pdf
# ---------------------------------------------------------------------------


@celery_app.task(
    name="app.workers.export_worker.export_to_pdf",
    queue="export-queue",
    soft_time_limit=300,
    time_limit=360,
)
def export_to_pdf(
    query_params: Dict[str, Any],
    output_path: str,
) -> Dict[str, Any]:
    """
    Export a summary PDF of matching businesses.
    """
    return _run_async(_export_to_pdf_async(query_params, output_path))


async def _export_to_pdf_async(
    query_params: Dict[str, Any],
    output_path: str,
) -> Dict[str, Any]:
    from reportlab.lib import colors  # type: ignore
    from reportlab.lib.pagesizes import A4, landscape  # type: ignore
    from reportlab.lib.styles import getSampleStyleSheet  # type: ignore
    from reportlab.lib.units import cm  # type: ignore
    from reportlab.platypus import (  # type: ignore
        Paragraph,
        SimpleDocTemplate,
        Spacer,
        Table,
        TableStyle,
    )

    businesses = await _query_businesses(query_params)
    out = _ensure_dir(output_path)

    doc = SimpleDocTemplate(
        str(out),
        pagesize=landscape(A4),
        rightMargin=1.5 * cm,
        leftMargin=1.5 * cm,
        topMargin=2 * cm,
        bottomMargin=2 * cm,
    )

    styles = getSampleStyleSheet()
    elements = []

    # Title
    elements.append(Paragraph("Business Intelligence Export", styles["Title"]))
    elements.append(
        Paragraph(
            f"Generated: {datetime.now(timezone.utc).strftime('%Y-%m-%d %H:%M UTC')} | "
            f"Total Records: {len(businesses)}",
            styles["Normal"],
        )
    )
    elements.append(Spacer(1, 0.5 * cm))

    # Table columns shown in PDF (subset to keep readable)
    pdf_cols = ["name", "city", "country", "website_url", "overall_score", "source"]
    headers = [col.replace("_", " ").title() for col in pdf_cols]

    table_data = [headers]
    for biz in businesses:
        row = [str(biz.get(col, "") or "") for col in pdf_cols]
        # Truncate long URLs
        url_idx = pdf_cols.index("website_url") if "website_url" in pdf_cols else -1
        if url_idx >= 0 and len(row[url_idx]) > 40:
            row[url_idx] = row[url_idx][:37] + "..."
        table_data.append(row)

    col_widths = [5 * cm, 3 * cm, 2.5 * cm, 6 * cm, 2.5 * cm, 3 * cm]

    table = Table(table_data, colWidths=col_widths, repeatRows=1)
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#2563EB")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("FONTSIZE", (0, 0), (-1, 0), 9),
                ("FONTSIZE", (0, 1), (-1, -1), 8),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F1F5F9")]),
                ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#CBD5E1")),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("TOPPADDING", (0, 0), (-1, -1), 4),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
            ]
        )
    )

    elements.append(table)
    doc.build(elements)

    file_size = out.stat().st_size
    logger.info(
        "export_pdf.done",
        output_path=output_path,
        row_count=len(businesses),
        file_size_bytes=file_size,
    )
    return {
        "output_path": output_path,
        "row_count": len(businesses),
        "file_size": file_size,
    }


# ---------------------------------------------------------------------------
# Task: generate_audit_pdf
# ---------------------------------------------------------------------------


@celery_app.task(
    name="app.workers.export_worker.generate_audit_pdf",
    queue="export-queue",
    soft_time_limit=300,
    time_limit=360,
)
def generate_audit_pdf(audit_id: str, output_path: str) -> Dict[str, Any]:
    """
    Generate a detailed PDF report for a single AuditReport.

    The report includes:
    - Business details
    - Score breakdown (radar-style table)
    - Recommendations
    - Tech stack
    - Performance metrics
    - AI analysis (if available)
    """
    return _run_async(_generate_audit_pdf_async(audit_id, output_path))


async def _generate_audit_pdf_async(
    audit_id: str, output_path: str
) -> Dict[str, Any]:
    from reportlab.lib import colors  # type: ignore
    from reportlab.lib.pagesizes import A4  # type: ignore
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet  # type: ignore
    from reportlab.lib.units import cm  # type: ignore
    from reportlab.platypus import (  # type: ignore
        HRFlowable,
        Paragraph,
        SimpleDocTemplate,
        Spacer,
        Table,
        TableStyle,
    )

    async with AsyncSessionLocal() as session:
        stmt = (
            select(AuditReport)
            .where(AuditReport.id == uuid.UUID(audit_id))
        )
        row = await session.execute(stmt)
        audit: Optional[AuditReport] = row.scalar_one_or_none()

        if audit is None:
            logger.error("generate_audit_pdf.not_found", audit_id=audit_id)
            return {"status": "not_found", "audit_id": audit_id}

        # Load related business
        biz_stmt = select(Business).where(Business.id == audit.business_id)
        biz_row = await session.execute(biz_stmt)
        business: Optional[Business] = biz_row.scalar_one_or_none()

    out = _ensure_dir(output_path)
    doc = SimpleDocTemplate(
        str(out),
        pagesize=A4,
        rightMargin=2 * cm,
        leftMargin=2 * cm,
        topMargin=2.5 * cm,
        bottomMargin=2.5 * cm,
    )

    styles = getSampleStyleSheet()
    brand_blue = colors.HexColor("#2563EB")

    heading_style = ParagraphStyle(
        "SectionHeading",
        parent=styles["Heading2"],
        textColor=brand_blue,
        spaceBefore=12,
        spaceAfter=6,
    )

    elements = []

    biz_name = business.name if business else "Unknown Business"
    biz_url = audit.website_url or (business.website_url if business else "N/A")

    # ---- Header ----
    elements.append(Paragraph(f"Website Audit Report", styles["Title"]))
    elements.append(Paragraph(biz_name, styles["Heading1"]))
    elements.append(
        Paragraph(
            f"URL: {biz_url}  |  "
            f"Generated: {datetime.now(timezone.utc).strftime('%Y-%m-%d %H:%M UTC')}  |  "
            f"Status: {audit.status.value.upper()}",
            styles["Normal"],
        )
    )
    elements.append(HRFlowable(width="100%", thickness=1, color=brand_blue))
    elements.append(Spacer(1, 0.5 * cm))

    # ---- Score summary table ----
    elements.append(Paragraph("Score Summary", heading_style))

    score_fields = [
        ("Overall Score", audit.overall_score),
        ("SEO", audit.seo_score),
        ("Mobile", audit.mobile_score),
        ("Performance", audit.speed_score),
        ("Accessibility", audit.accessibility_score),
        ("Security", audit.security_score),
        ("UI", audit.ui_score),
        ("UX", audit.ux_score),
        ("Content", audit.content_score),
        ("Conversion", audit.conversion_score),
        ("Branding", audit.branding_score),
    ]

    score_data = [["Dimension", "Score", "Rating"]]
    for label, score in score_fields:
        if score is None:
            continue
        rating = _score_rating(score)
        score_data.append([label, f"{score:.1f} / 100", rating])

    if len(score_data) > 1:
        score_table = Table(score_data, colWidths=[7 * cm, 4 * cm, 4 * cm])
        score_table.setStyle(
            TableStyle(
                [
                    ("BACKGROUND", (0, 0), (-1, 0), brand_blue),
                    ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                    ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                    ("FONTSIZE", (0, 0), (-1, -1), 9),
                    ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F1F5F9")]),
                    ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#CBD5E1")),
                    ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                    ("TOPPADDING", (0, 0), (-1, -1), 5),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
                ]
            )
        )
        elements.append(score_table)
        elements.append(Spacer(1, 0.5 * cm))

    # ---- Recommendations ----
    if audit.recommendations:
        elements.append(Paragraph("Recommendations", heading_style))
        for rec in audit.recommendations:
            priority = rec.get("priority", "medium").upper()
            title = rec.get("title", "")
            desc = rec.get("description", "")
            elements.append(
                Paragraph(
                    f"<b>[{priority}] {title}</b>: {desc}",
                    styles["Normal"],
                )
            )
            elements.append(Spacer(1, 0.2 * cm))

    # ---- Tech Stack ----
    if audit.tech_stack:
        elements.append(Paragraph("Technology Stack", heading_style))
        tech_items = [
            f"{t['name']} ({t['category']})"
            for t in audit.tech_stack
            if t.get("name")
        ]
        if tech_items:
            elements.append(Paragraph(", ".join(tech_items), styles["Normal"]))
        elements.append(Spacer(1, 0.3 * cm))

    # ---- Performance Metrics ----
    if audit.performance_metrics:
        elements.append(Paragraph("Performance Metrics", heading_style))
        perf_data = [["Metric", "Value"]]
        for key, val in audit.performance_metrics.items():
            if val is not None:
                perf_data.append([key.replace("_", " ").title(), str(val)])
        if len(perf_data) > 1:
            perf_table = Table(perf_data, colWidths=[7 * cm, 8 * cm])
            perf_table.setStyle(
                TableStyle(
                    [
                        ("BACKGROUND", (0, 0), (-1, 0), brand_blue),
                        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                        ("FONTSIZE", (0, 0), (-1, -1), 9),
                        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F1F5F9")]),
                        ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#CBD5E1")),
                        ("TOPPADDING", (0, 0), (-1, -1), 4),
                        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
                    ]
                )
            )
            elements.append(perf_table)
            elements.append(Spacer(1, 0.3 * cm))

    # ---- AI Analysis ----
    if audit.ai_analysis:
        ai = audit.ai_analysis
        elements.append(Paragraph("AI Analysis", heading_style))

        if ai.get("summary"):
            elements.append(Paragraph(ai["summary"], styles["Normal"]))
            elements.append(Spacer(1, 0.2 * cm))

        for section, label in [
            ("strengths", "Strengths"),
            ("weaknesses", "Weaknesses"),
            ("opportunities", "Opportunities"),
        ]:
            items = ai.get(section, [])
            if items:
                elements.append(Paragraph(f"<b>{label}:</b>", styles["Normal"]))
                for item in items:
                    elements.append(Paragraph(f"• {item}", styles["Normal"]))
                elements.append(Spacer(1, 0.1 * cm))

        if ai.get("estimated_value"):
            elements.append(
                Paragraph(
                    f"<b>Estimated Opportunity Value:</b> {ai['estimated_value'].upper()}",
                    styles["Normal"],
                )
            )

    doc.build(elements)

    file_size = out.stat().st_size
    logger.info(
        "generate_audit_pdf.done",
        audit_id=audit_id,
        output_path=output_path,
        file_size_bytes=file_size,
    )
    return {"output_path": output_path, "audit_id": audit_id, "file_size": file_size}


# ---------------------------------------------------------------------------
# Periodic: cleanup_old_exports
# ---------------------------------------------------------------------------


@celery_app.task(
    name="app.workers.export_worker.cleanup_old_exports",
    queue="export-queue",
)
def cleanup_old_exports() -> Dict[str, Any]:
    """
    Delete export files in the EXPORT_DIR that are older than 24 hours.
    """
    export_dir = Path(settings.EXPORT_DIR)
    if not export_dir.exists():
        return {"deleted": 0}

    cutoff = datetime.now(timezone.utc) - timedelta(hours=24)
    deleted = 0
    errors = 0

    for fpath in export_dir.rglob("*"):
        if not fpath.is_file():
            continue
        try:
            mtime = datetime.fromtimestamp(fpath.stat().st_mtime, tz=timezone.utc)
            if mtime < cutoff:
                fpath.unlink()
                deleted += 1
        except Exception as exc:
            errors += 1
            logger.warning("cleanup_exports.delete_error", path=str(fpath), error=str(exc))

    logger.info("cleanup_exports.done", deleted=deleted, errors=errors)
    return {"deleted": deleted, "errors": errors}


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------


def _score_rating(score: float) -> str:
    if score >= 90:
        return "Excellent"
    if score >= 75:
        return "Good"
    if score >= 60:
        return "Fair"
    if score >= 40:
        return "Poor"
    return "Critical"
