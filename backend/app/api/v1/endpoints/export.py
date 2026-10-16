"""
Export endpoints — synchronous generation, no Celery required.

Routes
------
POST /export/csv          – generate and download CSV immediately
POST /export/excel        – generate and download Excel immediately
POST /export/pdf          – generate and download PDF immediately
POST /export/jobs/csv     – async CSV (background), poll /export/jobs/{id}
POST /export/jobs/excel   – async Excel (background), poll /export/jobs/{id}
POST /export/jobs/pdf     – async PDF  (background), poll /export/jobs/{id}
GET  /export/jobs/{id}    – poll an async export job
GET  /export/download/{id} – download completed async export
"""
from __future__ import annotations

import asyncio
import csv
import html
import io
import logging
import os
import re
import time
import uuid
from datetime import datetime, timezone
from pathlib import Path
from typing import Annotated, Any, Dict, List, Optional

import httpx
from fastapi import APIRouter, BackgroundTasks, Depends, HTTPException, Query, status
from fastapi.responses import FileResponse, StreamingResponse
from pydantic import BaseModel, Field
from sqlalchemy import or_, select
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.config import settings
from app.core.database import get_db
from app.core.security import oauth2_scheme, verify_token
from app.models.business import Business
from app.models.discovery_result import DiscoveryResult
from app.models.user import User

logger = logging.getLogger(__name__)
router = APIRouter(prefix="/export", tags=["export"])

# In-memory job store for async exports  {job_id: {"status": ..., "path": ..., "error": ...}}
_export_jobs: Dict[str, Dict[str, Any]] = {}


# ---------------------------------------------------------------------------
# Website audit helper — lightweight, no external APIs
# ---------------------------------------------------------------------------

_SOCIAL_PATTERNS = {
    "Facebook": r'facebook\.com/(?!sharer)[^\s"\'<>?&]{3,60}',
    "Instagram": r'instagram\.com/[^\s"\'<>?&]{3,60}',
    "LinkedIn": r'linkedin\.com/(?:company|in)/[^\s"\'<>?&]{3,60}',
    "Twitter/X": r'(?:twitter|x)\.com/[^\s"\'<>?&]{3,60}',
    "TikTok": r'tiktok\.com/@[^\s"\'<>?&]{3,60}',
    "YouTube": r'youtube\.com/(?:channel|c|@)[^\s"\'<>?&]{3,60}',
    "WhatsApp": r'(?:wa\.me|whatsapp\.com/send)[^\s"\'<>?&]{0,60}',
}

_VIEWPORT_RE = re.compile(r'<meta[^>]+name=["\']viewport["\'][^>]+content', re.I)
_TITLE_RE = re.compile(r'<title[^>]*>([^<]{1,120})</title>', re.I)
_DESC_RE = re.compile(r'<meta[^>]+name=["\']description["\'][^>]+content=["\']([^"\']{1,300})', re.I)
_H1_RE = re.compile(r'<h1[^>]*>([^<]{1,120})</h1>', re.I)
_HTTPS_RE = re.compile(r'^https://', re.I)
_MAILTO_RE = re.compile(r'mailto:([a-zA-Z0-9._%+\-]+@[a-zA-Z0-9.\-]+\.[a-zA-Z]{2,})')
_TEL_RE = re.compile(r'tel:(\+?[\d\s\-().]{7,20})')
_SCHEMA_ORG_RE = re.compile(r'"@type"\s*:\s*"([A-Za-z]+)"')
_CANONICAL_RE = re.compile(r'<link[^>]+rel=["\']canonical["\'][^>]+href=["\']([^"\']+)', re.I)


async def _audit_website(url: str, timeout: float = 8.0) -> Dict[str, Any]:
    """
    Lightweight website audit: one HTTP request to homepage, parse HTML.
    Returns a dict with audit findings to embed in the PDF.
    """
    result: Dict[str, Any] = {
        "url": url,
        "reachable": False,
        "https": url.startswith("https://"),
        "status_code": None,
        "load_ms": None,
        "title": None,
        "meta_description": None,
        "h1": None,
        "canonical": None,
        "has_viewport": False,
        "has_schema_org": False,
        "emails": [],
        "phones": [],
        "social": {},
        "issues": [],
        "good": [],
    }
    try:
        hdrs = {
            "User-Agent": "Mozilla/5.0 (compatible; BI-Platform-Audit/1.0)",
            "Accept": "text/html,application/xhtml+xml",
        }
        t0 = time.monotonic()
        async with httpx.AsyncClient(
            follow_redirects=True, timeout=timeout, verify=False, headers=hdrs
        ) as client:
            resp = await client.get(url)
        result["load_ms"] = int((time.monotonic() - t0) * 1000)
        result["status_code"] = resp.status_code
        result["reachable"] = resp.status_code < 400

        if not result["reachable"]:
            result["issues"].append(f"Site returned HTTP {resp.status_code}")
            return result

        html = resp.text
        final_url = str(resp.url)

        # HTTPS check (after redirects)
        if not final_url.startswith("https://"):
            result["https"] = False
            result["issues"].append("Not using HTTPS — connection is insecure")
        else:
            result["https"] = True
            result["good"].append("HTTPS enabled")

        # Title
        m = _TITLE_RE.search(html)
        if m:
            result["title"] = m.group(1).strip()
            result["good"].append(f"Page title present: \"{result['title'][:60]}\"")
        else:
            result["issues"].append("Missing <title> tag — hurts SEO")

        # Meta description
        m = _DESC_RE.search(html)
        if m:
            result["meta_description"] = m.group(1).strip()
            result["good"].append("Meta description present")
        else:
            result["issues"].append("Missing meta description — hurts SEO click-through rate")

        # H1
        m = _H1_RE.search(html)
        if m:
            result["h1"] = m.group(1).strip()
            result["good"].append(f"H1 heading present: \"{result['h1'][:60]}\"")
        else:
            result["issues"].append("No <h1> heading found — hurts SEO")

        # Canonical URL
        m = _CANONICAL_RE.search(html)
        if m:
            result["canonical"] = m.group(1).strip()

        # Mobile viewport
        if _VIEWPORT_RE.search(html):
            result["has_viewport"] = True
            result["good"].append("Mobile viewport meta tag present")
        else:
            result["issues"].append("No viewport meta tag — site may not be mobile-friendly")

        # Schema.org structured data
        schema_types = list(set(_SCHEMA_ORG_RE.findall(html)))
        if schema_types:
            result["has_schema_org"] = True
            result["good"].append(f"Schema.org structured data: {', '.join(schema_types[:5])}")
        else:
            result["issues"].append("No Schema.org structured data — missed SEO opportunity")

        # Emails
        emails = list(set(_MAILTO_RE.findall(html)))
        result["emails"] = [e for e in emails if not re.match(r'^(noreply|no-reply|admin|example|test)@', e, re.I)][:3]

        # Phones
        phones_raw = list(set(_TEL_RE.findall(html)))
        result["phones"] = [re.sub(r'[\s]', '', p) for p in phones_raw[:3]]
        if result["phones"] or result["emails"]:
            result["good"].append("Contact info found on website")
        else:
            result["issues"].append("No contact info (phone/email) visible in page HTML")

        # Social media
        for platform, pattern in _SOCIAL_PATTERNS.items():
            found = re.search(pattern, html, re.I)
            if found:
                result["social"][platform] = found.group(0)

        if result["social"]:
            result["good"].append(f"Social profiles linked: {', '.join(result['social'].keys())}")
        else:
            result["issues"].append("No social media links found on homepage")

        # Load time rating
        ms = result["load_ms"]
        if ms and ms < 1500:
            result["good"].append(f"Fast page load: {ms}ms")
        elif ms and ms > 3000:
            result["issues"].append(f"Slow page load: {ms}ms — consider optimising")

    except httpx.ConnectError:
        result["issues"].append("Cannot connect to website — domain may be down or DNS failed")
    except httpx.TimeoutException:
        result["issues"].append(f"Website timed out after {timeout:.0f}s — very slow server")
    except Exception as exc:
        result["issues"].append(f"Audit error: {exc}")

    return result


async def _audit_websites_bulk(businesses: List[Any]) -> Dict[str, Dict[str, Any]]:
    """Run website audits for all businesses that have a URL, concurrently (max 8 at once)."""
    sem = asyncio.Semaphore(8)
    results: Dict[str, Dict[str, Any]] = {}

    async def _bounded(biz: Any) -> None:
        url = getattr(biz, "website_url", None)
        if not url:
            return
        async with sem:
            try:
                audit = await _audit_website(url)
                results[str(biz.id)] = audit
            except Exception as exc:
                logger.debug("Website audit failed for %s: %s", url, exc)

    await asyncio.gather(*[_bounded(b) for b in businesses])
    return results


# ---------------------------------------------------------------------------
# Opportunity analysis helper — project detection + scoring
# ---------------------------------------------------------------------------

_HIGH_VALUE_INDUSTRIES = {
    "construction", "real estate", "real_estate", "security", "hospital",
    "clinic", "medical", "dental", "healthcare", "education", "school",
    "university", "law", "legal", "lawyer", "accounting", "accountant",
    "finance", "logistics", "manufacturing", "factory", "gym", "fitness",
    "restaurant", "hotel", "resort", "hospitality", "travel", "agency",
    "retail", "ecommerce", "e-commerce", "automobile", "car dealer",
    "automotive", "import", "export", "corporate", "marketing",
}

_PROJECT_CATALOG = {
    "website_redesign": {
        "label": "Website Redesign",
        "value_usd": (3000, 15000),
        "complexity": "Medium",
        "trigger": "Outdated or missing website",
    },
    "mobile_app": {
        "label": "Mobile Application",
        "value_usd": (8000, 40000),
        "complexity": "High",
        "trigger": "No mobile app despite business need",
    },
    "crm_system": {
        "label": "CRM Implementation",
        "value_usd": (5000, 25000),
        "complexity": "High",
        "trigger": "Missing customer relationship management",
    },
    "booking_system": {
        "label": "Online Booking System",
        "value_usd": (2000, 8000),
        "complexity": "Medium",
        "trigger": "Service business without online booking",
    },
    "ai_chatbot": {
        "label": "AI Chatbot / WhatsApp Bot",
        "value_usd": (2500, 10000),
        "complexity": "Medium",
        "trigger": "High customer inquiry volume, no automation",
    },
    "seo_package": {
        "label": "SEO & Digital Marketing",
        "value_usd": (500, 3000),
        "complexity": "Low",
        "trigger": "Poor search visibility",
        "recurring": True,
    },
    "inventory_system": {
        "label": "Inventory Management System",
        "value_usd": (4000, 20000),
        "complexity": "High",
        "trigger": "Retail/manufacturing without digital inventory",
    },
    "erp_solution": {
        "label": "ERP / Business Management Platform",
        "value_usd": (15000, 80000),
        "complexity": "Very High",
        "trigger": "Mid-size company with manual processes",
    },
    "lead_management": {
        "label": "Lead Management System",
        "value_usd": (3000, 12000),
        "complexity": "Medium",
        "trigger": "Sales team without digital lead pipeline",
    },
    "payment_integration": {
        "label": "Online Payment / E-Commerce",
        "value_usd": (2000, 10000),
        "complexity": "Medium",
        "trigger": "Physical business without online sales",
    },
    "dashboard_analytics": {
        "label": "Business Dashboard & Analytics",
        "value_usd": (4000, 18000),
        "complexity": "High",
        "trigger": "No real-time business intelligence",
    },
    "pos_system": {
        "label": "POS / Order Management System",
        "value_usd": (3000, 15000),
        "complexity": "Medium",
        "trigger": "Restaurant/retail without modern POS",
    },
    "process_automation": {
        "label": "Process Automation",
        "value_usd": (5000, 30000),
        "complexity": "High",
        "trigger": "Manual repetitive workflows",
    },
}

# Industry → most relevant project types
_INDUSTRY_PROJECTS: Dict[str, List[str]] = {
    "construction":   ["website_redesign", "lead_management", "crm_system", "process_automation", "mobile_app"],
    "real_estate":    ["crm_system", "website_redesign", "lead_management", "mobile_app", "ai_chatbot"],
    "security":       ["crm_system", "website_redesign", "mobile_app", "lead_management", "process_automation"],
    "healthcare":     ["booking_system", "website_redesign", "crm_system", "mobile_app", "ai_chatbot"],
    "medical":        ["booking_system", "website_redesign", "crm_system", "mobile_app", "ai_chatbot"],
    "dental":         ["booking_system", "website_redesign", "ai_chatbot", "crm_system"],
    "clinic":         ["booking_system", "website_redesign", "crm_system", "ai_chatbot"],
    "hospital":       ["booking_system", "erp_solution", "mobile_app", "crm_system", "process_automation"],
    "education":      ["website_redesign", "erp_solution", "mobile_app", "crm_system", "ai_chatbot"],
    "school":         ["website_redesign", "mobile_app", "crm_system", "ai_chatbot"],
    "legal":          ["crm_system", "website_redesign", "lead_management", "process_automation", "ai_chatbot"],
    "law":            ["crm_system", "website_redesign", "lead_management", "ai_chatbot"],
    "accounting":     ["crm_system", "process_automation", "website_redesign", "erp_solution"],
    "finance":        ["crm_system", "process_automation", "erp_solution", "website_redesign"],
    "logistics":      ["erp_solution", "mobile_app", "process_automation", "dashboard_analytics", "crm_system"],
    "manufacturing":  ["erp_solution", "inventory_system", "process_automation", "dashboard_analytics"],
    "fitness":        ["booking_system", "mobile_app", "crm_system", "website_redesign"],
    "gym":            ["booking_system", "mobile_app", "crm_system", "website_redesign"],
    "restaurant":     ["pos_system", "website_redesign", "ai_chatbot", "mobile_app", "payment_integration"],
    "hotel":          ["booking_system", "crm_system", "website_redesign", "mobile_app", "ai_chatbot"],
    "resort":         ["booking_system", "website_redesign", "crm_system", "mobile_app"],
    "travel":         ["booking_system", "crm_system", "website_redesign", "mobile_app", "ai_chatbot"],
    "ecommerce":      ["website_redesign", "mobile_app", "crm_system", "ai_chatbot", "dashboard_analytics"],
    "retail":         ["pos_system", "inventory_system", "website_redesign", "mobile_app", "payment_integration"],
    "automotive":     ["crm_system", "website_redesign", "lead_management", "booking_system", "ai_chatbot"],
    "marketing":      ["crm_system", "lead_management", "ai_chatbot", "dashboard_analytics", "process_automation"],
    "default":        ["website_redesign", "seo_package", "crm_system", "lead_management", "ai_chatbot"],
}


def _classify_industry(biz: Any) -> str:
    """Return a normalized industry key matching _INDUSTRY_PROJECTS."""
    raw = (
        (getattr(biz, "industry", None) or "") + " " +
        (getattr(biz, "category", None) or "")
    ).lower()
    for key in _INDUSTRY_PROJECTS:
        if key == "default":
            continue
        if key in raw:
            return key
    return "default"


def _compute_opportunity_score(biz: Any, wa: Optional[Dict[str, Any]]) -> int:
    """
    0–100 opportunity score.
    Higher = more gaps to fill = bigger sales opportunity.
    """
    score = 40  # base

    # +10 high-value industry
    industry_raw = ((getattr(biz, "industry", None) or "") + " " + (getattr(biz, "category", None) or "")).lower()
    if any(k in industry_raw for k in _HIGH_VALUE_INDUSTRIES):
        score += 10

    # Contact availability signals
    if getattr(biz, "phone", None):
        score += 5
    if getattr(biz, "email", None):
        score += 5

    # Website gaps (big opportunity)
    if not getattr(biz, "has_website", True):
        score += 20  # no website = huge gap
    elif wa:
        if not wa.get("https"):
            score += 5
        if not wa.get("has_viewport"):
            score += 5
        if not wa.get("has_schema_org"):
            score += 3
        if not wa.get("meta_description"):
            score += 3
        issues_count = len(wa.get("issues") or [])
        score += min(issues_count * 2, 10)
        load_ms = wa.get("load_ms") or 0
        if load_ms > 3000:
            score += 5
        if not wa.get("social"):
            score += 4

    # Social media presence (partial = opportunity for growth)
    social_count = sum(1 for f in [
        getattr(biz, "facebook_url", None), getattr(biz, "instagram_url", None),
        getattr(biz, "linkedin_url", None), getattr(biz, "twitter_url", None),
    ] if f)
    if social_count == 0:
        score += 8
    elif social_count < 3:
        score += 3

    # Completeness — lower completeness = more unknown = more research needed
    cs = getattr(biz, "completeness_score", None) or 0
    if cs < 40:
        score += 5

    return min(score, 100)


def _get_project_recommendations(biz: Any, wa: Optional[Dict[str, Any]]) -> List[Dict[str, Any]]:
    """Return top 3–5 project recommendations for this business."""
    industry_key = _classify_industry(biz)
    project_keys = _INDUSTRY_PROJECTS.get(industry_key, _INDUSTRY_PROJECTS["default"])[:5]

    recs = []
    for pk in project_keys:
        cat = _PROJECT_CATALOG.get(pk)
        if not cat:
            continue
        lo, hi = cat["value_usd"]
        recs.append({
            "type": cat["label"],
            "complexity": cat["complexity"],
            "value_min": lo,
            "value_max": hi,
            "trigger": cat["trigger"],
            "recurring": cat.get("recurring", False),
        })

    # Boost website-related recommendations if audit found issues
    if wa and wa.get("reachable"):
        issues = wa.get("issues") or []
        has_mobile_issue = any("mobile" in i.lower() or "viewport" in i.lower() for i in issues)
        has_seo_issue = any("seo" in i.lower() or "title" in i.lower() or "description" in i.lower() for i in issues)
        has_speed_issue = any("slow" in i.lower() or "load" in i.lower() for i in issues)

        extras = []
        if has_mobile_issue and not any(r["type"] == "Website Redesign" for r in recs):
            extras.append(_PROJECT_CATALOG["website_redesign"])
        if has_seo_issue and not any(r["type"] == "SEO & Digital Marketing" for r in recs):
            extras.append(_PROJECT_CATALOG["seo_package"])
        for cat in extras:
            lo, hi = cat["value_usd"]
            recs.insert(0, {
                "type": cat["label"],
                "complexity": cat["complexity"],
                "value_min": lo,
                "value_max": hi,
                "trigger": cat["trigger"],
                "recurring": cat.get("recurring", False),
            })

    # Deduplicate
    seen = set()
    deduped = []
    for r in recs:
        if r["type"] not in seen:
            seen.add(r["type"])
            deduped.append(r)
    return deduped[:5]


def _total_value_range(recs: List[Dict[str, Any]]) -> tuple:
    lo = sum(r["value_min"] for r in recs)
    hi = sum(r["value_max"] for r in recs)
    return lo, hi


# ---------------------------------------------------------------------------
# Auth
# ---------------------------------------------------------------------------

async def _get_current_user(
    token: Annotated[str, Depends(oauth2_scheme)],
    db: AsyncSession = Depends(get_db),
) -> User:
    err = HTTPException(
        status_code=status.HTTP_401_UNAUTHORIZED,
        detail="Could not validate credentials",
        headers={"WWW-Authenticate": "Bearer"},
    )
    try:
        token_data = verify_token(token, expected_type="access")
    except Exception:
        raise err
    result = await db.execute(select(User).where(User.id == uuid.UUID(token_data.sub)))
    user = result.scalar_one_or_none()
    if not user or not user.is_active:
        raise err
    return user


# ---------------------------------------------------------------------------
# Pydantic schemas
# ---------------------------------------------------------------------------

class ExportFilters(BaseModel):
    industry: Optional[str] = None
    country: Optional[str] = None
    state: Optional[str] = None
    city: Optional[str] = None
    has_website: Optional[bool] = None
    min_score: Optional[float] = Field(None, ge=0, le=100)
    max_score: Optional[float] = Field(None, ge=0, le=100)
    source: Optional[str] = None
    discovery_job_id: Optional[str] = Field(None, description="Filter to businesses from a specific discovery job")
    ids: Optional[List[str]] = Field(None, description="Explicit list of business UUIDs")
    limit: int = Field(1000, ge=1, le=100000)


# ---------------------------------------------------------------------------
# Business query helper
# ---------------------------------------------------------------------------

async def _query_businesses(filters: ExportFilters, db: AsyncSession) -> List[Business]:
    stmt = select(Business)

    if filters.ids:
        try:
            uuids = [uuid.UUID(i) for i in filters.ids]
            stmt = stmt.where(Business.id.in_(uuids))
        except ValueError:
            pass
    else:
        if filters.industry:
            stmt = stmt.where(Business.industry == filters.industry)
        if filters.country:
            stmt = stmt.where(Business.country == filters.country)
        if filters.state:
            stmt = stmt.where(Business.state == filters.state)
        if filters.city:
            stmt = stmt.where(Business.city == filters.city)
        if filters.has_website is not None:
            if filters.has_website:
                stmt = stmt.where(
                    or_(
                        Business.has_website == True,
                        (Business.website_url.isnot(None)) & (Business.website_url != ""),
                    )
                )
            else:
                stmt = stmt.where(
                    Business.has_website == False,
                    or_(Business.website_url.is_(None), Business.website_url == ""),
                )
        if filters.min_score is not None:
            stmt = stmt.where(Business.overall_score >= filters.min_score)
        if filters.max_score is not None:
            stmt = stmt.where(Business.overall_score <= filters.max_score)
        if filters.source:
            stmt = stmt.where(Business.source == filters.source)
        if filters.discovery_job_id:
            try:
                job_uuid = uuid.UUID(filters.discovery_job_id)
                # Check if DiscoveryResult records exist for this job
                dr_count_row = await db.execute(
                    select(DiscoveryResult.business_id)
                    .where(DiscoveryResult.job_id == job_uuid)
                    .where(DiscoveryResult.business_id.is_not(None))
                    .limit(1)
                )
                has_results = dr_count_row.first() is not None

                if has_results:
                    # Modern path: exact match via DiscoveryResult
                    biz_ids_subq = (
                        select(DiscoveryResult.business_id)
                        .where(DiscoveryResult.job_id == job_uuid)
                        .where(DiscoveryResult.business_id.is_not(None))
                        .scalar_subquery()
                    )
                    stmt = stmt.where(Business.id.in_(biz_ids_subq))
                else:
                    # Legacy fallback: filter by job's city/country from config
                    from app.models.discovery_job import DiscoveryJob
                    job_row = await db.execute(
                        select(DiscoveryJob).where(DiscoveryJob.id == job_uuid)
                    )
                    job = job_row.scalar_one_or_none()
                    if job and job.config:
                        cfg = job.config
                        _city = cfg.get("city")
                        _country = cfg.get("country")
                        _src_list = cfg.get("sources", [])
                        # Filter by sources used in the job
                        if _src_list:
                            stmt = stmt.where(Business.source.in_(_src_list))
                        # Also narrow by city/country if specified
                        if _city:
                            stmt = stmt.where(Business.city.ilike(f"%{_city}%"))
                        elif _country:
                            stmt = stmt.where(Business.country.ilike(f"%{_country}%"))
                        # Limit to businesses discovered around job start/end time
                        if job.started_at:
                            stmt = stmt.where(Business.discovered_at >= job.started_at)
                        if job.completed_at:
                            stmt = stmt.where(Business.discovered_at <= job.completed_at)
            except ValueError:
                pass

    stmt = stmt.order_by(Business.completeness_score.desc().nulls_last(), Business.created_at.desc())
    stmt = stmt.limit(filters.limit)
    rows = await db.execute(stmt)
    return list(rows.scalars().all())


# ---------------------------------------------------------------------------
# CSV generation
# ---------------------------------------------------------------------------

_CSV_COLUMNS = [
    "name", "category", "industry", "description",
    "address", "city", "state", "country", "zip_code",
    "latitude", "longitude", "google_maps_url",
    "phone", "phone2", "whatsapp", "email", "email2",
    "website_url", "contact_page_url", "about_page_url",
    "facebook_url", "instagram_url", "linkedin_url", "twitter_url",
    "tiktok_url", "youtube_url", "pinterest_url",
    "rating", "review_count", "founded_year", "company_size",
    "overall_score", "opportunity_score", "completeness_score",
    "source", "discovered_at",
]


def _businesses_to_csv(businesses: List[Business]) -> str:
    output = io.StringIO()
    writer = csv.DictWriter(output, fieldnames=_CSV_COLUMNS, extrasaction="ignore")
    writer.writeheader()
    for b in businesses:
        row: Dict[str, Any] = {}
        for col in _CSV_COLUMNS:
            val = getattr(b, col if col != "metadata_" else "metadata_", None)
            if isinstance(val, datetime):
                val = val.isoformat()
            row[col] = val if val is not None else ""
        writer.writerow(row)
    return output.getvalue()


# ---------------------------------------------------------------------------
# Excel generation
# ---------------------------------------------------------------------------

def _businesses_to_excel(businesses: List[Business]) -> bytes:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Businesses"

    header_fill = PatternFill(start_color="1E40AF", end_color="1E40AF", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True)

    for col_idx, col_name in enumerate(_CSV_COLUMNS, 1):
        cell = ws.cell(row=1, column=col_idx, value=col_name.replace("_", " ").title())
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")
        ws.column_dimensions[ws.cell(row=1, column=col_idx).column_letter].width = max(12, len(col_name) + 2)

    for row_idx, b in enumerate(businesses, 2):
        for col_idx, col in enumerate(_CSV_COLUMNS, 1):
            val = getattr(b, col, None)
            if isinstance(val, datetime):
                val = val.isoformat()
            elif isinstance(val, (list, dict)):
                import json
                val = json.dumps(val)
            ws.cell(row=row_idx, column=col_idx, value=val)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------------------
# PDF generation
# ---------------------------------------------------------------------------

def _businesses_to_pdf(
    businesses: List[Business],
    job_name: str = "Business Intelligence Report",
    web_audits: Optional[Dict[str, Dict[str, Any]]] = None,
) -> bytes:
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import cm
    from reportlab.lib import colors
    from reportlab.platypus import (
        SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle,
        HRFlowable, KeepTogether, PageBreak,
    )
    from reportlab.lib.enums import TA_LEFT, TA_CENTER

    buf = io.BytesIO()
    doc = SimpleDocTemplate(
        buf,
        pagesize=A4,
        rightMargin=1.5 * cm,
        leftMargin=1.5 * cm,
        topMargin=2 * cm,
        bottomMargin=2 * cm,
    )

    styles = getSampleStyleSheet()
    brand_blue = colors.HexColor("#1E40AF")
    light_gray = colors.HexColor("#F1F5F9")
    dark_gray = colors.HexColor("#374151")
    green = colors.HexColor("#16A34A")
    yellow = colors.HexColor("#D97706")
    red = colors.HexColor("#DC2626")

    style_title = ParagraphStyle(
        "ReportTitle", parent=styles["Title"],
        textColor=brand_blue, fontSize=22, spaceAfter=6,
    )
    style_subtitle = ParagraphStyle(
        "Subtitle", parent=styles["Normal"],
        textColor=dark_gray, fontSize=10, spaceAfter=12,
    )
    style_section = ParagraphStyle(
        "Section", parent=styles["Heading2"],
        textColor=brand_blue, fontSize=12, spaceBefore=10, spaceAfter=4,
    )
    style_biz_name = ParagraphStyle(
        "BizName", parent=styles["Heading1"],
        textColor=brand_blue, fontSize=16, spaceAfter=4,
    )
    style_body = ParagraphStyle(
        "Body", parent=styles["Normal"],
        fontSize=9, textColor=dark_gray, leading=14,
    )
    style_link = ParagraphStyle(
        "Link", parent=style_body,
        textColor=brand_blue,
    )

    # Pre-compute opportunity scores and sort businesses descending
    _opp_cache: Dict[str, int] = {}
    for _b in businesses:
        _wa = (web_audits or {}).get(str(_b.id))
        _opp_cache[str(_b.id)] = _compute_opportunity_score(_b, _wa)
    businesses = sorted(businesses, key=lambda b: _opp_cache.get(str(b.id), 0), reverse=True)

    story = []

    # ── Cover page ──
    story.append(Spacer(1, 2 * cm))
    story.append(Paragraph(html.escape(job_name), style_title))
    story.append(Paragraph(
        f"Generated: {datetime.now(timezone.utc).strftime('%B %d, %Y')} &nbsp;&nbsp;|&nbsp;&nbsp; "
        f"{len(businesses)} high-value lead{'s' if len(businesses) != 1 else ''} &nbsp;&nbsp;|&nbsp;&nbsp; "
        f"Sorted by Opportunity Score",
        style_subtitle,
    ))
    story.append(HRFlowable(width="100%", thickness=2, color=brand_blue))
    story.append(Spacer(1, 0.4 * cm))

    # Summary stats
    complete_80 = sum(1 for b in businesses if b.completeness_score and b.completeness_score >= 80)
    with_website = sum(1 for b in businesses if b.has_website)
    with_email = sum(1 for b in businesses if b.email)
    with_phone = sum(1 for b in businesses if b.phone)
    with_social = sum(1 for b in businesses if any([b.facebook_url, b.instagram_url, b.linkedin_url, b.twitter_url]))
    audited_count = len(web_audits) if web_audits else 0
    online_count = sum(1 for a in (web_audits or {}).values() if a.get("reachable"))
    excellent = sum(1 for s in _opp_cache.values() if s >= 80)
    high_opp = sum(1 for s in _opp_cache.values() if 60 <= s < 80)
    all_recs = []
    for _b in businesses:
        _wa = (web_audits or {}).get(str(_b.id))
        all_recs.extend(_get_project_recommendations(_b, _wa))
    total_min = sum(r["value_min"] for r in all_recs)
    total_max = sum(r["value_max"] for r in all_recs)

    summary_data = [
        ["Metric", "Value"],
        ["Total Leads", str(len(businesses))],
        ["Excellent Opportunity (80+)", str(excellent)],
        ["High Opportunity (60–79)", str(high_opp)],
        ["Have Website", str(with_website)],
        ["Websites Audited", str(audited_count)],
        ["Websites Online", str(online_count)],
        ["Have Email", str(with_email)],
        ["Have Phone", str(with_phone)],
        ["Have Social Media", str(with_social)],
        ["≥80% Profile Complete", str(complete_80)],
        ["Total Project Value (Min)", f"${total_min:,.0f}"],
        ["Total Project Value (Max)", f"${total_max:,.0f}"],
    ]
    summary_table = Table(summary_data, colWidths=[10 * cm, 4 * cm])
    summary_table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), brand_blue),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 10),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, light_gray]),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.lightgrey),
        ("PADDING", (0, 0), (-1, -1), 6),
    ]))
    story.append(summary_table)
    story.append(Spacer(1, 0.5 * cm))

    # Top 10 leads table on cover
    story.append(Paragraph("<b>Top Leads by Opportunity Score</b>", style_section))
    top_rows = [["#", "Business", "Industry", "City", "Opp. Score", "Est. Value"]]
    for _i, _b in enumerate(businesses[:10], 1):
        _wa = (web_audits or {}).get(str(_b.id))
        _recs = _get_project_recommendations(_b, _wa)
        _lo, _hi = _total_value_range(_recs)
        _score = _opp_cache.get(str(_b.id), 0)
        _score_color = "#16A34A" if _score >= 80 else ("#D97706" if _score >= 60 else "#DC2626")
        top_rows.append([
            str(_i),
            Paragraph(f"<b>{html.escape((_b.name or '')[:35])}</b>", style_body),
            Paragraph(html.escape((_b.industry or _b.category or "—")[:25]), style_body),
            Paragraph(html.escape((_b.city or "—")[:20]), style_body),
            Paragraph(f'<font color="{_score_color}"><b>{_score}</b></font>', style_body),
            Paragraph(f"${_lo:,.0f}–${_hi:,.0f}", style_body),
        ])
    top_table = Table(top_rows, colWidths=[0.7*cm, 5*cm, 3.2*cm, 2.5*cm, 1.8*cm, 3.3*cm])
    top_table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), brand_blue),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 8),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, light_gray]),
        ("GRID", (0, 0), (-1, -1), 0.3, colors.lightgrey),
        ("PADDING", (0, 0), (-1, -1), 4),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("ALIGN", (0, 0), (0, -1), "CENTER"),
        ("ALIGN", (4, 0), (4, -1), "CENTER"),
    ]))
    story.append(top_table)
    story.append(PageBreak())

    # ── One page per business ──
    for biz in businesses:
        blocks = []

        # Header row: name + opportunity score badge + completeness badge
        cs = biz.completeness_score or 0
        badge_color = green if cs >= 80 else (yellow if cs >= 50 else red)
        _h_opp = _opp_cache.get(str(biz.id), _compute_opportunity_score(biz, (web_audits or {}).get(str(biz.id))))
        opp_badge_color = green if _h_opp >= 80 else (yellow if _h_opp >= 60 else red)
        name_para = Paragraph(f"<b>{html.escape(biz.name or 'Unknown Business')}</b>", style_biz_name)

        opp_badge_table = Table([[f"Opp: {_h_opp}/100"]], colWidths=[2.5 * cm])
        opp_badge_table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, -1), opp_badge_color),
            ("TEXTCOLOR", (0, 0), (-1, -1), colors.white),
            ("FONTNAME", (0, 0), (-1, -1), "Helvetica-Bold"),
            ("FONTSIZE", (0, 0), (-1, -1), 9),
            ("ALIGN", (0, 0), (-1, -1), "CENTER"),
            ("PADDING", (0, 0), (-1, -1), 4),
        ]))
        badge_table = Table([[f"{cs:.0f}% Complete"]], colWidths=[2.5 * cm])
        badge_table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, -1), badge_color),
            ("TEXTCOLOR", (0, 0), (-1, -1), colors.white),
            ("FONTNAME", (0, 0), (-1, -1), "Helvetica-Bold"),
            ("FONTSIZE", (0, 0), (-1, -1), 9),
            ("ALIGN", (0, 0), (-1, -1), "CENTER"),
            ("PADDING", (0, 0), (-1, -1), 4),
        ]))

        header_table = Table([[name_para, opp_badge_table, badge_table]], colWidths=[11.5 * cm, 2.7 * cm, 2.3 * cm])
        header_table.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "MIDDLE")]))
        blocks.append(header_table)

        if biz.category or biz.industry:
            blocks.append(Paragraph(
                f"<i>{html.escape(biz.category or biz.industry)}</i>",
                ParagraphStyle("Cat", parent=style_body, textColor=colors.grey, fontSize=9),
            ))
        blocks.append(HRFlowable(width="100%", thickness=1, color=colors.lightgrey))
        blocks.append(Spacer(1, 0.2 * cm))

        # Description
        if biz.description:
            blocks.append(Paragraph("<b>About</b>", style_section))
            blocks.append(Paragraph(html.escape(biz.description[:600]), style_body))
            blocks.append(Spacer(1, 0.2 * cm))

        # Contact details table
        contact_rows: List[List[str]] = []

        def _add(label: str, val: Optional[str], link: bool = False) -> None:
            if val:
                if link:
                    contact_rows.append([label, Paragraph(f'<a href="{html.escape(val)}" color="blue">{html.escape(val)}</a>', style_link)])
                else:
                    contact_rows.append([label, Paragraph(html.escape(val), style_body)])

        _add("Address", biz.address)
        # Build city/state/country line
        _loc_parts = [p for p in [biz.city, biz.state, biz.zip_code, biz.country] if p]
        _add("Location", ", ".join(_loc_parts) if _loc_parts else None)
        # Coordinates
        if biz.latitude and biz.longitude:
            contact_rows.append(["Coordinates", Paragraph(f"{biz.latitude:.5f}, {biz.longitude:.5f}", style_body)])
        # Google Maps link — prefer stored URL, fall back to coords
        _maps_url = biz.google_maps_url
        if not _maps_url and biz.latitude and biz.longitude:
            _maps_url = f"https://www.google.com/maps?q={biz.latitude},{biz.longitude}"
        elif not _maps_url and biz.address:
            import urllib.parse
            _maps_url = f"https://www.google.com/maps/search/?api=1&query={urllib.parse.quote(biz.address + ' ' + (biz.city or '') + ' ' + (biz.country or ''))}"
        _add("Google Maps", _maps_url, link=True)
        _add("Phone", biz.phone)
        _add("Phone 2", biz.phone2)
        _add("WhatsApp", biz.whatsapp)
        _add("Email", biz.email, link=True)
        _add("Email 2", biz.email2, link=True)
        _add("Website", biz.website_url, link=True)
        _add("Contact Page", biz.contact_page_url, link=True)
        if biz.rating:
            contact_rows.append(["Rating", Paragraph(f"{biz.rating:.1f} / 5 ({biz.review_count or 0} reviews)", style_body)])
        if biz.founded_year:
            contact_rows.append(["Founded", Paragraph(str(biz.founded_year), style_body)])
        if biz.company_size:
            contact_rows.append(["Company Size", Paragraph(biz.company_size, style_body)])

        if contact_rows:
            blocks.append(Paragraph("<b>Contact & Location</b>", style_section))
            ct = Table(contact_rows, colWidths=[3.5 * cm, 13 * cm])
            ct.setStyle(TableStyle([
                ("FONTNAME", (0, 0), (0, -1), "Helvetica-Bold"),
                ("FONTSIZE", (0, 0), (-1, -1), 9),
                ("ROWBACKGROUNDS", (0, 0), (-1, -1), [colors.white, light_gray]),
                ("GRID", (0, 0), (-1, -1), 0.3, colors.lightgrey),
                ("PADDING", (0, 0), (-1, -1), 5),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ]))
            blocks.append(ct)
            blocks.append(Spacer(1, 0.3 * cm))

        # Social media
        social_items = []
        if biz.facebook_url:
            social_items.append(("Facebook", biz.facebook_url))
        if biz.instagram_url:
            social_items.append(("Instagram", biz.instagram_url))
        if biz.linkedin_url:
            social_items.append(("LinkedIn", biz.linkedin_url))
        if biz.twitter_url:
            social_items.append(("X / Twitter", biz.twitter_url))
        if biz.tiktok_url:
            social_items.append(("TikTok", biz.tiktok_url))
        if biz.youtube_url:
            social_items.append(("YouTube", biz.youtube_url))
        if biz.pinterest_url:
            social_items.append(("Pinterest", biz.pinterest_url))

        if social_items:
            blocks.append(Paragraph("<b>Social Media</b>", style_section))
            sm_rows = [
                [label, Paragraph(f'<a href="{url}" color="blue">{url}</a>', style_link)]
                for label, url in social_items
            ]
            sm_table = Table(sm_rows, colWidths=[3.5 * cm, 13 * cm])
            sm_table.setStyle(TableStyle([
                ("FONTNAME", (0, 0), (0, -1), "Helvetica-Bold"),
                ("FONTSIZE", (0, 0), (-1, -1), 9),
                ("ROWBACKGROUNDS", (0, 0), (-1, -1), [colors.white, light_gray]),
                ("GRID", (0, 0), (-1, -1), 0.3, colors.lightgrey),
                ("PADDING", (0, 0), (-1, -1), 5),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ]))
            blocks.append(sm_table)
            blocks.append(Spacer(1, 0.3 * cm))

        # Scores
        score_data = []
        if biz.opportunity_score:
            score_data.append(["Opportunity Score", f"{biz.opportunity_score:.1f}"])
        if biz.overall_score:
            score_data.append(["Overall Score", f"{biz.overall_score:.1f}"])
        if biz.revenue_potential:
            score_data.append(["Revenue Potential", f"${biz.revenue_potential:,.0f}"])
        if biz.discovered_at:
            score_data.append(["Discovered", biz.discovered_at.strftime("%Y-%m-%d")])
        score_data.append(["Source", biz.source or "unknown"])

        if score_data:
            blocks.append(Paragraph("<b>Intelligence Scores</b>", style_section))
            sc_table = Table(score_data, colWidths=[5 * cm, 11.5 * cm])
            sc_table.setStyle(TableStyle([
                ("FONTNAME", (0, 0), (0, -1), "Helvetica-Bold"),
                ("FONTSIZE", (0, 0), (-1, -1), 9),
                ("ROWBACKGROUNDS", (0, 0), (-1, -1), [colors.white, light_gray]),
                ("GRID", (0, 0), (-1, -1), 0.3, colors.lightgrey),
                ("PADDING", (0, 0), (-1, -1), 5),
            ]))
            blocks.append(sc_table)
            blocks.append(Spacer(1, 0.3 * cm))

        # ── Opportunity Analysis ──
        blocks.append(HRFlowable(width="100%", thickness=1, color=colors.HexColor("#CBD5E1"), spaceBefore=4))
        blocks.append(Paragraph("<b>Opportunity Analysis</b>", style_section))

        _biz_wa = (web_audits or {}).get(str(biz.id))
        _opp_score = _compute_opportunity_score(biz, _biz_wa)
        _opp_color = "#16A34A" if _opp_score >= 80 else ("#D97706" if _opp_score >= 60 else "#DC2626")
        _opp_label = "Excellent" if _opp_score >= 80 else ("High" if _opp_score >= 60 else ("Medium" if _opp_score >= 40 else "Low"))
        _recs = _get_project_recommendations(biz, _biz_wa)
        _total_lo, _total_hi = _total_value_range(_recs)

        # Opportunity score pill
        opp_summary_rows = [
            ["Opportunity Score", Paragraph(f'<font color="{_opp_color}"><b>{_opp_score}/100 — {_opp_label}</b></font>', style_body)],
            ["Industry", Paragraph((biz.industry or biz.category or "Unknown").title(), style_body)],
            ["Total Est. Project Value", Paragraph(f"<b>${_total_lo:,.0f} – ${_total_hi:,.0f} USD</b>", style_body)],
        ]
        # Gaps detected
        _gaps = []
        if not biz.has_website:
            _gaps.append("No website")
        if not biz.email:
            _gaps.append("No email found")
        if not biz.phone:
            _gaps.append("No phone found")
        if not any([biz.facebook_url, biz.instagram_url, biz.linkedin_url, biz.twitter_url]):
            _gaps.append("No social media presence")
        if _biz_wa:
            for _issue in (_biz_wa.get("issues") or [])[:4]:
                _gaps.append(html.escape(_issue))
        if _gaps:
            opp_summary_rows.append(["Gaps Identified", Paragraph(" • ".join(_gaps[:6]), style_body)])

        opp_sum_tbl = Table(opp_summary_rows, colWidths=[4.5*cm, 12*cm])
        opp_sum_tbl.setStyle(TableStyle([
            ("FONTNAME", (0, 0), (0, -1), "Helvetica-Bold"),
            ("FONTSIZE", (0, 0), (-1, -1), 9),
            ("ROWBACKGROUNDS", (0, 0), (-1, -1), [colors.white, light_gray]),
            ("GRID", (0, 0), (-1, -1), 0.3, colors.lightgrey),
            ("PADDING", (0, 0), (-1, -1), 5),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ]))
        blocks.append(opp_sum_tbl)
        blocks.append(Spacer(1, 0.3 * cm))

        # Project recommendations table
        if _recs:
            blocks.append(Paragraph(
                "<b>Recommended Services &amp; Projects</b>",
                ParagraphStyle("ProjHead", parent=style_body, fontSize=9, fontName="Helvetica-Bold", spaceBefore=4, spaceAfter=3),
            ))
            proj_rows = [["Project / Service", "Complexity", "Est. Value (USD)", "Why"]]
            for _r in _recs:
                _recurring = " /mo" if _r.get("recurring") else ""
                proj_rows.append([
                    Paragraph(f"<b>{html.escape(_r['type'])}</b>", style_body),
                    Paragraph(html.escape(_r["complexity"]), style_body),
                    Paragraph(f"${_r['value_min']:,.0f}–${_r['value_max']:,.0f}{_recurring}", style_body),
                    Paragraph(html.escape(_r["trigger"]), ParagraphStyle("PTrigger", parent=style_body, fontSize=8)),
                ])
            proj_tbl = Table(proj_rows, colWidths=[4.0*cm, 2.0*cm, 3.5*cm, 7.0*cm])
            proj_tbl.setStyle(TableStyle([
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1E40AF")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("FONTSIZE", (0, 0), (-1, -1), 8),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, light_gray]),
                ("GRID", (0, 0), (-1, -1), 0.3, colors.lightgrey),
                ("PADDING", (0, 0), (-1, -1), 5),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ]))
            blocks.append(proj_tbl)
            blocks.append(Spacer(1, 0.3 * cm))

        # Website audit section
        wa = (web_audits or {}).get(str(biz.id))
        if wa:
            blocks.append(HRFlowable(width="100%", thickness=1, color=colors.HexColor("#CBD5E1"), spaceBefore=6))
            blocks.append(Paragraph("<b>Website Audit</b>", style_section))

            # Summary row
            _status_color = green if wa.get("reachable") else red
            _status_text = "Online" if wa.get("reachable") else "Offline / Unreachable"
            wa_summary_rows = [
                ["URL", Paragraph(f'<a href="{wa["url"]}" color="blue">{wa["url"]}</a>', style_link)],
                ["Status", Paragraph(f'<font color="{"#16A34A" if wa.get("reachable") else "#DC2626"}"><b>{_status_text}</b></font>', style_body)],
            ]
            if wa.get("status_code"):
                wa_summary_rows.append(["HTTP Code", Paragraph(str(wa["status_code"]), style_body)])
            if wa.get("https") is not None:
                _https_txt = '<font color="#16A34A"><b>Yes ✓</b></font>' if wa["https"] else '<font color="#DC2626"><b>No ✗</b></font>'
                wa_summary_rows.append(["HTTPS", Paragraph(_https_txt, style_body)])
            if wa.get("load_ms"):
                _ms = wa["load_ms"]
                _ms_color = "#16A34A" if _ms < 1500 else ("#D97706" if _ms < 3000 else "#DC2626")
                wa_summary_rows.append(["Load Time", Paragraph(f'<font color="{_ms_color}"><b>{_ms} ms</b></font>', style_body)])
            if wa.get("title"):
                wa_summary_rows.append(["Page Title", Paragraph(html.escape(wa["title"][:80]), style_body)])
            if wa.get("meta_description"):
                wa_summary_rows.append(["Meta Desc", Paragraph(html.escape(wa["meta_description"][:120]), style_body)])
            if wa.get("h1"):
                wa_summary_rows.append(["H1 Heading", Paragraph(html.escape(wa["h1"][:80]), style_body)])

            wa_sum_table = Table(wa_summary_rows, colWidths=[3.5 * cm, 13 * cm])
            wa_sum_table.setStyle(TableStyle([
                ("FONTNAME", (0, 0), (0, -1), "Helvetica-Bold"),
                ("FONTSIZE", (0, 0), (-1, -1), 9),
                ("ROWBACKGROUNDS", (0, 0), (-1, -1), [colors.white, light_gray]),
                ("GRID", (0, 0), (-1, -1), 0.3, colors.lightgrey),
                ("PADDING", (0, 0), (-1, -1), 5),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ]))
            blocks.append(wa_sum_table)
            blocks.append(Spacer(1, 0.25 * cm))

            # Contact info found on website (may differ from business record)
            wa_phones = wa.get("phones") or []
            wa_emails = wa.get("emails") or []
            if wa_phones or wa_emails:
                blocks.append(Paragraph(
                    "<b>Contact Found on Website</b>",
                    ParagraphStyle("WASubHead", parent=style_body, fontSize=9, fontName="Helvetica-Bold", spaceBefore=4, spaceAfter=2),
                ))
                contact_found_rows = []
                for ph in wa_phones:
                    contact_found_rows.append(["Phone", Paragraph(ph, style_body)])
                for em in wa_emails:
                    contact_found_rows.append(["Email", Paragraph(f'<a href="mailto:{em}" color="blue">{em}</a>', style_link)])
                cf_table = Table(contact_found_rows, colWidths=[3.5 * cm, 13 * cm])
                cf_table.setStyle(TableStyle([
                    ("FONTNAME", (0, 0), (0, -1), "Helvetica-Bold"),
                    ("FONTSIZE", (0, 0), (-1, -1), 9),
                    ("ROWBACKGROUNDS", (0, 0), (-1, -1), [colors.white, light_gray]),
                    ("GRID", (0, 0), (-1, -1), 0.3, colors.lightgrey),
                    ("PADDING", (0, 0), (-1, -1), 5),
                ]))
                blocks.append(cf_table)
                blocks.append(Spacer(1, 0.2 * cm))

            # Social profiles found on website
            wa_social = wa.get("social") or {}
            if wa_social:
                blocks.append(Paragraph(
                    "<b>Social Profiles Found on Website</b>",
                    ParagraphStyle("WASubHead2", parent=style_body, fontSize=9, fontName="Helvetica-Bold", spaceBefore=4, spaceAfter=2),
                ))
                soc_rows = [
                    [platform, Paragraph(f'<a href="https://{path}" color="blue">https://{path}</a>', style_link)]
                    for platform, path in wa_social.items()
                ]
                soc_table = Table(soc_rows, colWidths=[3.5 * cm, 13 * cm])
                soc_table.setStyle(TableStyle([
                    ("FONTNAME", (0, 0), (0, -1), "Helvetica-Bold"),
                    ("FONTSIZE", (0, 0), (-1, -1), 9),
                    ("ROWBACKGROUNDS", (0, 0), (-1, -1), [colors.white, light_gray]),
                    ("GRID", (0, 0), (-1, -1), 0.3, colors.lightgrey),
                    ("PADDING", (0, 0), (-1, -1), 5),
                ]))
                blocks.append(soc_table)
                blocks.append(Spacer(1, 0.2 * cm))

            # Issues & recommendations
            issues = wa.get("issues") or []
            good = wa.get("good") or []
            if good:
                blocks.append(Paragraph(
                    "<b>What's Working</b>",
                    ParagraphStyle("WAGoodHead", parent=style_body, fontSize=9, fontName="Helvetica-Bold", spaceBefore=4, spaceAfter=2),
                ))
                for item in good:
                    blocks.append(Paragraph(
                        f'<font color="#16A34A">✓</font> {html.escape(item)}',
                        ParagraphStyle("WAGood", parent=style_body, leftIndent=12, spaceAfter=1),
                    ))
                blocks.append(Spacer(1, 0.15 * cm))
            if issues:
                blocks.append(Paragraph(
                    "<b>Issues &amp; Recommendations</b>",
                    ParagraphStyle("WAIssueHead", parent=style_body, fontSize=9, fontName="Helvetica-Bold", spaceBefore=4, spaceAfter=2),
                ))
                for item in issues:
                    blocks.append(Paragraph(
                        f'<font color="#DC2626">✗</font> {html.escape(item)}',
                        ParagraphStyle("WAIssue", parent=style_body, leftIndent=12, spaceAfter=1),
                    ))
                blocks.append(Spacer(1, 0.2 * cm))

        elif biz.website_url:
            # Had a URL but no audit result (shouldn't happen normally)
            blocks.append(Spacer(1, 0.2 * cm))
            blocks.append(Paragraph(
                f'<i>Website: <a href="{biz.website_url}" color="blue">{biz.website_url}</a> — audit not available</i>',
                ParagraphStyle("NoAudit", parent=style_body, textColor=colors.grey),
            ))

        story.append(KeepTogether(blocks))
        story.append(PageBreak())

    def _on_page(canvas, doc):
        canvas.saveState()
        canvas.setFont("Helvetica", 8)
        canvas.setFillColor(colors.grey)
        canvas.drawString(1.5 * cm, 1.2 * cm, "BI Platform — Business Intelligence Report")
        canvas.drawRightString(A4[0] - 1.5 * cm, 1.2 * cm, f"Page {doc.page}")
        canvas.restoreState()

    doc.build(story, onFirstPage=_on_page, onLaterPages=_on_page)
    return buf.getvalue()


# ---------------------------------------------------------------------------
# Output path helper
# ---------------------------------------------------------------------------

def _build_output_path(job_id: str, extension: str) -> str:
    export_dir = Path(getattr(settings, "EXPORT_DIR", "/tmp/bi_exports"))
    export_dir.mkdir(parents=True, exist_ok=True)
    return str(export_dir / f"export_{job_id}{extension}")


# ---------------------------------------------------------------------------
# Synchronous endpoints (immediate download)
# ---------------------------------------------------------------------------

@router.post("/csv")
async def export_csv_sync(
    filters: ExportFilters = ExportFilters(),
    db: AsyncSession = Depends(get_db),
    current_user: Any = Depends(_get_current_user),
) -> StreamingResponse:
    """Generate and stream a CSV export immediately."""
    businesses = await _query_businesses(filters, db)
    csv_text = _businesses_to_csv(businesses)
    return StreamingResponse(
        io.StringIO(csv_text),
        media_type="text/csv",
        headers={"Content-Disposition": f"attachment; filename=export_{uuid.uuid4().hex[:8]}.csv"},
    )


@router.post("/excel")
async def export_excel_sync(
    filters: ExportFilters = ExportFilters(),
    db: AsyncSession = Depends(get_db),
    current_user: Any = Depends(_get_current_user),
) -> StreamingResponse:
    """Generate and stream an Excel export immediately."""
    import asyncio
    businesses = await _query_businesses(filters, db)
    loop = asyncio.get_event_loop()
    xlsx_bytes = await loop.run_in_executor(None, _businesses_to_excel, businesses)
    return StreamingResponse(
        io.BytesIO(xlsx_bytes),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename=export_{uuid.uuid4().hex[:8]}.xlsx"},
    )


@router.post("/pdf")
async def export_pdf_sync(
    filters: ExportFilters = ExportFilters(),
    job_name: str = Query("Business Intelligence Report"),
    db: AsyncSession = Depends(get_db),
    current_user: Any = Depends(_get_current_user),
) -> StreamingResponse:
    """Generate and stream a PDF export with live website audits embedded."""
    businesses = await _query_businesses(filters, db)
    # Audit all websites concurrently before building PDF
    web_audits = await _audit_websites_bulk(businesses)
    loop = asyncio.get_event_loop()
    pdf_bytes = await loop.run_in_executor(None, _businesses_to_pdf, businesses, job_name, web_audits)
    return StreamingResponse(
        io.BytesIO(pdf_bytes),
        media_type="application/pdf",
        headers={"Content-Disposition": f"attachment; filename=report_{uuid.uuid4().hex[:8]}.pdf"},
    )


# ---------------------------------------------------------------------------
# Async (background) endpoints — for large exports
# ---------------------------------------------------------------------------

async def _run_async_export(job_id: str, fmt: str, filters: ExportFilters, job_name: str) -> None:
    """Background task: generate export file, update _export_jobs."""
    from app.core.database import AsyncSessionLocal
    import asyncio

    _export_jobs[job_id] = {"status": "running", "format": fmt, "path": None, "error": None}
    try:
        async with AsyncSessionLocal() as db:
            businesses = await _query_businesses(filters, db)

        loop = asyncio.get_event_loop()
        if fmt == "csv":
            content = _businesses_to_csv(businesses).encode("utf-8")
            ext = ".csv"
        elif fmt == "excel":
            content = await loop.run_in_executor(None, _businesses_to_excel, businesses)
            ext = ".xlsx"
        else:
            web_audits = await _audit_websites_bulk(businesses)
            content = await loop.run_in_executor(None, _businesses_to_pdf, businesses, job_name, web_audits)
            ext = ".pdf"

        path = _build_output_path(job_id, ext)
        with open(path, "wb") as f:
            f.write(content if isinstance(content, bytes) else content)
        _export_jobs[job_id] = {"status": "done", "format": fmt, "path": path, "error": None, "count": len(businesses)}
    except Exception as exc:
        logger.exception("Async export failed: %s", exc)
        _export_jobs[job_id] = {"status": "failed", "format": fmt, "path": None, "error": str(exc)}


@router.post("/jobs/csv", status_code=status.HTTP_202_ACCEPTED)
async def export_csv_async(
    background_tasks: BackgroundTasks,
    filters: ExportFilters = ExportFilters(),
    current_user: Any = Depends(_get_current_user),
) -> Dict[str, Any]:
    job_id = uuid.uuid4().hex
    background_tasks.add_task(_run_async_export, job_id, "csv", filters, "Export")
    return {"job_id": job_id, "format": "csv", "status": "queued"}


@router.post("/jobs/excel", status_code=status.HTTP_202_ACCEPTED)
async def export_excel_async(
    background_tasks: BackgroundTasks,
    filters: ExportFilters = ExportFilters(),
    current_user: Any = Depends(_get_current_user),
) -> Dict[str, Any]:
    job_id = uuid.uuid4().hex
    background_tasks.add_task(_run_async_export, job_id, "excel", filters, "Export")
    return {"job_id": job_id, "format": "excel", "status": "queued"}


@router.post("/jobs/pdf", status_code=status.HTTP_202_ACCEPTED)
async def export_pdf_async(
    background_tasks: BackgroundTasks,
    filters: ExportFilters = ExportFilters(),
    job_name: str = Query("Business Intelligence Report"),
    current_user: Any = Depends(_get_current_user),
) -> Dict[str, Any]:
    job_id = uuid.uuid4().hex
    background_tasks.add_task(_run_async_export, job_id, "pdf", filters, job_name)
    return {"job_id": job_id, "format": "pdf", "status": "queued"}


@router.get("/jobs/{job_id}")
async def get_export_job_status(
    job_id: str,
    current_user: Any = Depends(_get_current_user),
) -> Dict[str, Any]:
    job = _export_jobs.get(job_id)
    if not job:
        raise HTTPException(status_code=404, detail="Export job not found")
    return {"job_id": job_id, **job}


@router.get("/download/{job_id}")
async def download_export(
    job_id: str,
    current_user: Any = Depends(_get_current_user),
) -> FileResponse:
    job = _export_jobs.get(job_id)
    if not job or job["status"] != "done":
        raise HTTPException(
            status_code=404,
            detail=f"Export not ready (status: {job['status'] if job else 'not found'})",
        )
    path = job["path"]
    if not path or not os.path.exists(path):
        raise HTTPException(status_code=404, detail="Export file not found on disk")

    ext = Path(path).suffix.lower()
    media_map = {
        ".csv": "text/csv",
        ".xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        ".pdf": "application/pdf",
    }
    return FileResponse(
        path=path,
        media_type=media_map.get(ext, "application/octet-stream"),
        filename=f"export_{job_id}{ext}",
    )
